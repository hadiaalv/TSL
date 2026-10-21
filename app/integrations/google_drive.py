"""Read-only Google Drive search for the authenticated CoS account."""

from __future__ import annotations

import csv
import difflib
import io
import logging
import re
from urllib.parse import quote
from html.parser import HTMLParser
import httplib2
from typing import Any, Optional

logger = logging.getLogger(__name__)

_DOC_MIME_TYPE = "application/vnd.google-apps.document"
_SHEET_MIME_TYPE = "application/vnd.google-apps.spreadsheet"
_SLIDE_MIME_TYPE = "application/vnd.google-apps.presentation"
_FOLDER_MIME_TYPE = "application/vnd.google-apps.folder"
_MAX_CONTENT_CHARS = 6000
_SHEET_MAX_ROWS = 500
_SHEET_MAX_COLUMNS = 50
_SEARCH_TERM_VARIANTS = {
    "cubsi": ("Cubsi", "CUBSY"),
    "cubsy": ("Cubsi", "CUBSY"),
}


class _HTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        if tag.lower() in ("script", "style", "noscript"):
            self._ignored_depth += 1
        elif tag.lower() in ("p", "div", "br", "li", "tr", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in ("script", "style", "noscript") and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag.lower() in ("p", "div", "li", "tr", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth and data.strip():
            self.parts.append(data.strip())


def _html_text(content: str) -> str:
    parser = _HTMLTextParser()
    parser.feed(content)
    return "\n".join(part for part in parser.parts if part.strip())


def _extract_downloaded_file(content: bytes, mime_type: str) -> tuple[str, bool]:
    """Extract searchable text from common uploaded document formats."""
    if mime_type == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(page for page in pages if page.strip())
        if not text.strip():
            raise ValueError(
                "This PDF contains no selectable text; scanned PDFs require OCR."
            )
        return text, False

    if mime_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        from docx import Document

        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        paragraphs.extend(
            "\n".join(
                "\t".join(cell.text for cell in row.cells)
                for row in table.rows
            )
            for table in document.tables
        )
        return "\n\n".join(paragraphs), False

    if mime_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sections = []
        truncated = False
        for worksheet in workbook.worksheets:
            row_limit = min(worksheet.max_row or 1, _SHEET_MAX_ROWS)
            column_limit = min(worksheet.max_column or 1, _SHEET_MAX_COLUMNS)
            truncated = truncated or (worksheet.max_row or 1) > row_limit
            truncated = truncated or (worksheet.max_column or 1) > column_limit
            rows = worksheet.iter_rows(
                max_row=row_limit,
                max_col=column_limit,
                values_only=True,
            )
            text_rows = [
                "\t".join("" if value is None else str(value) for value in row)
                for row in rows
            ]
            sections.append(f"Sheet: {worksheet.title}\n" + "\n".join(text_rows))
        workbook.close()
        return "\n\n".join(sections), truncated

    if mime_type == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ):
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
        slides = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text.strip())
                if shape.has_table:
                    parts.extend(
                        "\t".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
            slides.append(f"Slide {slide_number}\n" + "\n".join(parts))
        return "\n\n".join(slides), False

    if mime_type in ("text/csv", "application/csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text))), False

    if mime_type == "text/html":
        return _html_text(content.decode("utf-8", errors="replace")), False

    if mime_type.startswith("text/") or mime_type in (
        "application/json",
        "application/xml",
        "application/javascript",
        "application/x-yaml",
    ):
        return content.decode("utf-8", errors="replace"), False

    raise TypeError(
        "Text extraction is not available for this file type "
        f"({mime_type or 'unknown type'})."
    )


def _column_name(column_number: int) -> str:
    """Return the A1 notation column name for a one-based column number."""
    name = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        name = chr(ord("A") + remainder) + name
    return name


def _spreadsheet_text(
    spreadsheets_service: Any,
    spreadsheet_id: str,
    operations: list[str],
) -> tuple[str, bool]:
    operations.append("sheets.spreadsheets.get")
    spreadsheet = (
        spreadsheets_service.spreadsheets()
        .get(
            spreadsheetId=spreadsheet_id,
            fields="sheets(properties(title,gridProperties(rowCount,columnCount)))",
        )
        .execute(num_retries=3)
    )
    sheet_properties = [
        sheet.get("properties", {})
        for sheet in spreadsheet.get("sheets", [])
    ]
    ranges = []
    truncated = False
    for properties in sheet_properties:
        title = properties.get("title", "")
        grid = properties.get("gridProperties", {})
        available_rows = grid.get("rowCount", 1)
        available_columns = grid.get("columnCount", 1)
        row_count = min(available_rows, _SHEET_MAX_ROWS)
        column_count = min(available_columns, _SHEET_MAX_COLUMNS)
        truncated = truncated or available_rows > row_count
        truncated = truncated or available_columns > column_count
        escaped_title = title.replace("'", "''")
        ranges.append(
            f"'{escaped_title}'!A1:{_column_name(column_count)}{row_count}"
        )

    if not ranges:
        return "", truncated

    operations.append("sheets.spreadsheets.values.batchGet")
    result = (
        spreadsheets_service.spreadsheets()
        .values()
        .batchGet(
            spreadsheetId=spreadsheet_id,
            ranges=ranges,
            valueRenderOption="FORMATTED_VALUE",
        )
        .execute(num_retries=3)
    )
    value_ranges = result.get("valueRanges", [])
    sections = []
    for properties, value_range in zip(sheet_properties, value_ranges):
        title = properties.get("title", "(untitled sheet)")
        rows = value_range.get("values", [])
        if not rows:
            sections.append(f"Sheet: {title}\n(empty)")
            continue
        text_rows = ["\t".join(str(value) for value in row) for row in rows]
        sections.append(f"Sheet: {title}\n" + "\n".join(text_rows))
    return "\n\n".join(sections), truncated


def _safe_error_message(exc: Exception) -> str:
    message = str(getattr(exc, "reason", None) or exc)
    message = re.sub(r"(?i)\bBearer\s+\S+", "Bearer [redacted]", message)
    message = re.sub(
        r"(?i)\b(access[_ -]?token|refresh[_ -]?token|client[_ -]?secret|authorization|cookie)\b"
        r"\s*[:=]\s*\S+",
        r"\1=[redacted]",
        message,
    )
    return message[:500]


def _presentation_text(presentations_service: Any, presentation_id: str) -> str:
    presentation = (
        presentations_service.presentations()
        .get(
            presentationId=presentation_id,
            fields="slides(pageElements(shape(text(textElements(textRun(content))))))",
        )
        .execute(num_retries=3)
    )
    paragraphs = []
    for slide in presentation.get("slides", []):
        for element in slide.get("pageElements", []):
            shape = element.get("shape", {})
            text = shape.get("text", {})
            for text_element in text.get("textElements", []):
                run = text_element.get("textRun", {})
                if run.get("content"):
                    paragraphs.append(run["content"])
    return "\n".join(paragraphs)


def _read_file_content(
    service: Any,
    credentials: Any,
    item: dict[str, Any],
) -> dict[str, Any]:
    from googleapiclient.discovery import build
    from googleapiclient.errors import HttpError

    mime_type = item.get("mimeType", "")
    export_type = {_DOC_MIME_TYPE: "text/plain"}.get(mime_type)
    extractable_download_types = {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "text/csv",
        "application/csv",
        "text/html",
        "application/json",
        "application/xml",
        "application/javascript",
        "application/x-yaml",
    }
    if not (
        export_type
        or mime_type in (_SHEET_MIME_TYPE, _SLIDE_MIME_TYPE, "text/plain")
        or mime_type.startswith("text/")
        or mime_type in extractable_download_types
    ):
        return {
            "content_available": False,
            "content_status": "unsupported",
            "content_error": (
                "Text extraction is not available for this file type "
                f"({mime_type or 'unknown type'})."
            ),
        }

    operations = []
    content_truncated_by_format = False
    operation = "drive.files.export" if export_type else "drive.files.get_media"
    if mime_type == _SHEET_MIME_TYPE:
        operation = "sheets API client initialization"
    elif mime_type == _SLIDE_MIME_TYPE:
        operation = "slides API client initialization"
    try:
        if export_type:
            content = (
                service.files()
                .export(fileId=item["id"], mimeType=export_type)
                .execute(num_retries=3)
            )
            if isinstance(content, bytes):
                content = content.decode("utf-8", errors="replace")
            content_format = export_type
        elif mime_type == _SHEET_MIME_TYPE:
            spreadsheets_service = build(
                "sheets",
                "v4",
                credentials=credentials,
                cache_discovery=False,
            )
            content, content_truncated_by_format = _spreadsheet_text(
                spreadsheets_service,
                item["id"],
                operations,
            )
            content_format = "text/plain"
        elif mime_type == _SLIDE_MIME_TYPE:
            slides_service = build(
                "slides",
                "v1",
                credentials=credentials,
                cache_discovery=False,
            )
            content = _presentation_text(slides_service, item["id"])
            content_format = "text/plain"
        else:
            downloaded = service.files().get_media(fileId=item["id"]).execute(
                num_retries=3
            )
            if not isinstance(downloaded, bytes):
                raise TypeError("Drive returned file data in an unsupported format.")
            content, content_truncated_by_format = _extract_downloaded_file(
                downloaded,
                mime_type,
            )
            content_format = "text/plain"
        if not isinstance(content, str):
            raise TypeError("Drive returned content in an unsupported format.")
        if not content.strip():
            raise ValueError("This file is empty or contains no extractable text.")
        return {
            "content": content[:_MAX_CONTENT_CHARS],
            "content_format": content_format,
            "content_truncated": (
                len(content) > _MAX_CONTENT_CHARS or content_truncated_by_format
            ),
            "content_available": True,
            "content_status": "read",
        }
    except Exception as exc:  # Keep a single unreadable file from aborting search.
        if operations:
            operation = " -> ".join(operations)
        status_code = getattr(getattr(exc, "resp", None), "status", None)
        error_message = _safe_error_message(exc)
        logger.warning(
            "Drive content read failed file_name=%r file_id=%s "
            "mime_type=%s operation=%s http_status=%s "
            "google_api_error=%r exception_type=%s",
            item.get("name", "(untitled)"),
            item.get("id"),
            mime_type,
            operation,
            status_code,
            error_message,
            type(exc).__name__,
        )
        if isinstance(exc, ModuleNotFoundError):
            missing_module = exc.name or "document parser"
            error_message = (
                f"Required parser dependency '{missing_module}' is not installed. "
                "Install project dependencies with 'pip install -r requirements.txt'."
            )
        elif (
            isinstance(exc, (HttpError, httplib2.HttpLib2Error, OSError))
            or status_code is not None
        ):
            error_message = _file_read_error(status_code)
        else:
            error_message = _safe_error_message(exc)
        return {
            "content_available": False,
            "content_status": "error",
            "content_error": error_message,
        }


def _file_read_error(status_code: Any) -> str:
    if status_code == 401:
        return "Google authorization expired. Reconnect the Google account."
    if status_code == 403:
        return "The connected Google account does not have permission to read this file."
    if status_code == 404:
        return "This Drive file was deleted or is no longer accessible."
    if status_code == 429:
        return "Google Drive rate-limited this file after retries. Try again later."
    if status_code and status_code >= 500:
        return "Google Drive had a temporary error reading this file. Try again later."
    return "Google could not read this file. Try again or check its file permissions."


def verify_drive_access(user_id: str) -> dict[str, str]:
    """Verify credentials and read-only Drive API access with a real API call."""
    from googleapiclient.discovery import build
    from googleapiclient.errors import HttpError

    from app import db
    from app.integrations.google_services import (
        DRIVE_READ_SCOPE,
        GoogleConnectionError,
        load_credentials,
    )

    row = db.get_google_token(user_id)
    if not row or not (row.get("token") or row.get("refresh_token")):
        return {
            "drive_status": "not_connected",
            "drive_error": "Connect your Google account to check Drive access.",
        }

    if DRIVE_READ_SCOPE not in (row.get("scopes") or []):
        return {
            "drive_status": "permission_missing",
            "drive_error": (
                "Google Drive read-only permission is missing. Reconnect Google "
                "and approve Drive access."
            ),
        }

    try:
        credentials = load_credentials(user_id)
    except GoogleConnectionError as exc:
        return {
            "drive_status": (
                "reconnect_required"
                if exc.requires_reconnect
                else "verification_failed"
            ),
            "drive_error": str(exc),
        }
    if credentials is None:
        return {
            "drive_status": "not_connected",
            "drive_error": "Connect your Google account to check Drive access.",
        }

    service = build("drive", "v3", credentials=credentials, cache_discovery=False)
    try:
        service.files().list(
            q="trashed = false",
            pageSize=1,
            fields="files(id)",
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        ).execute(num_retries=3)
    except (HttpError, httplib2.HttpLib2Error, OSError) as exc:
        status_code = getattr(getattr(exc, "resp", None), "status", None)
        if status_code == 401:
            message = "Google authorization is no longer valid. Reconnect Google."
            db.set_google_connection_state(
                user_id,
                "reconnect_required",
                message,
            )
            return {
                "drive_status": "reconnect_required",
                "drive_error": message,
            }
        if status_code == 403:
            return {
                "drive_status": "permission_missing",
                "drive_error": (
                    "Google Drive access was denied. Reconnect and approve the "
                    "Drive read-only permission; confirm the Drive API is enabled "
                    "in the configured Google Cloud project."
                ),
            }
        logger.warning(
            "Drive connectivity check failed user_id=%s http_status=%s",
            user_id,
            status_code,
        )
        return {
            "drive_status": "verification_failed",
            "drive_error": (
                "Google Drive could not be checked right now. Try again shortly."
            ),
        }
    except (httplib2.HttpLib2Error, OSError) as exc:
        logger.warning(
            "Drive connectivity check unavailable user_id=%s exception_type=%s",
            user_id,
            type(exc).__name__,
        )
        return {
            "drive_status": "verification_failed",
            "drive_error": (
                "Google Drive could not be reached right now. Try again shortly."
            ),
        }
    db.set_google_connection_state(user_id, "connected")
    return {"drive_status": "verified", "drive_error": ""}


def fetch_recent_drive_files(
    user_id: str,
    search_text: str = "",
    max_results: int = 5,
) -> list[dict[str, Any]]:
    from googleapiclient.discovery import build
    from googleapiclient.errors import HttpError

    from app.integrations.google_services import load_credentials

    credentials = load_credentials(user_id)
    if not credentials:
        raise RuntimeError("Google not connected. Connect Google from the CoS account first.")
    terms = re.findall(r"[A-Za-z0-9_-]{1,80}", search_text)[:5]
    search_terms = []
    for term in terms:
        search_terms.extend(_SEARCH_TERM_VARIANTS.get(term.casefold(), (term,)))
    search_terms = list(dict.fromkeys(search_terms))
    query = "trashed = false"
    if search_terms:
        term_query = " or ".join(
            f"fullText contains '{term}'" for term in search_terms
        )
        query += f" and ({term_query})"

    service = build("drive", "v3", credentials=credentials, cache_discovery=False)
    result = (
        service.files()
        .list(
            q=query,
            orderBy="modifiedTime desc",
            pageSize=max(1, min(max_results, 10)),
            fields=(
                "files(id,name,mimeType,createdTime,modifiedTime,description,"
                "webViewLink,parents)"
            ),
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        )
        .execute(num_retries=3)
    )

    parent_cache: dict[str, dict[str, Any]] = {}

    def get_parent(parent_id: str) -> dict[str, Any]:
        if parent_id in parent_cache:
            return parent_cache[parent_id]
        try:
            parent_cache[parent_id] = (
                service.files()
                .get(
                    fileId=parent_id,
                    fields="id,name,mimeType,webViewLink",
                    supportsAllDrives=True,
                )
                .execute(num_retries=3)
            )
        except HttpError as exc:
            message = _safe_error_message(exc)
            status_code = getattr(getattr(exc, "resp", None), "status", None)
            logger.warning(
                "Drive parent lookup failed parent_id=%s http_status=%s "
                "exception_type=%s",
                parent_id,
                status_code,
                type(exc).__name__,
            )
            parent_cache[parent_id] = {
                "id": parent_id,
                "lookup_error": message,
            }
        return parent_cache[parent_id]

    files: list[dict[str, Any]] = []
    for item in result.get("files", []):
        parent_ids = item.get("parents", [])
        file_data = {
            "id": item.get("id"),
            "name": item.get("name", "(untitled)"),
            "mime_type": item.get("mimeType", ""),
            "created_time": item.get("createdTime"),
            "modified_time": item.get("modifiedTime"),
            "description": item.get("description", ""),
            "web_view_link": item.get("webViewLink"),
            "parent_ids": parent_ids,
            "parents": [get_parent(parent_id) for parent_id in parent_ids],
        }
        if file_data["mime_type"] == _FOLDER_MIME_TYPE:
            file_data["is_folder"] = True
            file_data["content_available"] = False
            file_data["content_status"] = "folder"
            files.append(file_data)
            continue

        file_data.update(_read_file_content(service, credentials, item))
        files.append(file_data)

    return files


def fetch_brand_drive_files(
    user_id: str,
    businesses: list[dict[str, Any]],
    portfolio_folder_id: Optional[str] = None,
    folder_mappings: Optional[dict[str, str]] = None,
    file_mappings: Optional[dict[str, list[str]]] = None,
) -> dict[str, dict[str, Any]]:
    """Link files from uniquely matched brand folders beneath a verified root."""
    from googleapiclient.discovery import build
    from googleapiclient.errors import HttpError

    if not portfolio_folder_id:
        return {
            str(business["id"]): {
                "status": "portfolio_not_configured",
                "business_name": business.get("name", ""),
                "files": [],
                "error": (
                    "Set GOOGLE_DRIVE_PORTFOLIO_FOLDER_ID to the verified "
                    "portfolio folder before syncing brand folders."
                ),
            }
            for business in businesses
        }

    from app.integrations.google_services import load_credentials

    credentials = load_credentials(user_id)
    if not credentials:
        raise RuntimeError("Google not connected. Connect Google from the CoS account first.")
    service = build("drive", "v3", credentials=credentials, cache_discovery=False)
    portfolio_root = (
        service.files()
        .get(
            fileId=portfolio_folder_id,
            fields="id,name,mimeType",
            supportsAllDrives=True,
        )
        .execute(num_retries=3)
    )
    if portfolio_root.get("mimeType") != _FOLDER_MIME_TYPE:
        raise RuntimeError(
            "GOOGLE_DRIVE_PORTFOLIO_FOLDER_ID does not identify a Google Drive folder."
        )

    folders: list[dict[str, Any]] = []
    page_token = None
    while True:
        params = {
            "q": f"trashed = false and mimeType = '{_FOLDER_MIME_TYPE}'",
            "pageSize": 1000,
            "fields": "nextPageToken,files(id,name,mimeType,webViewLink,parents)",
            "supportsAllDrives": True,
            "includeItemsFromAllDrives": True,
            "corpora": "allDrives",
        }
        if page_token:
            params["pageToken"] = page_token
        result = service.files().list(**params).execute(num_retries=3)
        folders.extend(result.get("files", []))
        page_token = result.get("nextPageToken")
        if not page_token:
            break

    def folder_key(value: str) -> str:
        return "".join(character for character in value.casefold() if character.isalnum())

    children: dict[str, list[str]] = {}
    for folder in folders:
        for parent_id in folder.get("parents", []):
            children.setdefault(parent_id, []).append(folder["id"])

    portfolio_descendants: set[str] = {portfolio_folder_id}
    pending = [portfolio_folder_id]
    while pending:
        parent_id = pending.pop()
        for child_id in children.get(parent_id, []):
            if child_id not in portfolio_descendants:
                portfolio_descendants.add(child_id)
                pending.append(child_id)

    folder_index: dict[str, list[dict[str, Any]]] = {}
    folders_by_id = {folder["id"]: folder for folder in folders if folder.get("id")}
    for folder in folders:
        if (
            folder.get("id") == portfolio_folder_id
            or folder.get("id") not in portfolio_descendants
        ):
            continue
        folder_index.setdefault(folder_key(folder.get("name", "")), []).append(folder)

    def folder_details(folder: dict[str, Any]) -> dict[str, str]:
        names = [str(folder.get("name") or "(unnamed folder)")]
        parent_ids = folder.get("parents", [])
        parent_id = parent_ids[0] if parent_ids else None
        visited = {folder.get("id")}
        while parent_id and parent_id not in visited:
            visited.add(parent_id)
            parent = folders_by_id.get(parent_id)
            if parent:
                names.append(str(parent.get("name") or "(unnamed folder)"))
                parent_ids = parent.get("parents", [])
                parent_id = parent_ids[0] if parent_ids else None
                continue
            if parent_id == portfolio_folder_id:
                names.append(str(portfolio_root.get("name") or "Portfolio"))
            break
        return {
            "id": str(folder.get("id") or ""),
            "name": str(folder.get("name") or ""),
            "path": " / ".join(reversed(names)),
            "web_view_link": str(folder.get("webViewLink") or ""),
        }

    matches: dict[str, dict[str, Any]] = {}
    for business in businesses:
        business_id = str(business["id"])
        business_name = str(business.get("name") or "")
        business_key = folder_key(str(business.get("name") or ""))
        candidates = folder_index.get(business_key, [])
        mapped_folder_id = (folder_mappings or {}).get(business_id) or (
            folder_mappings or {}
        ).get(business_name)
        mapped_file_ids = list(
            dict.fromkeys(
                (file_mappings or {}).get(business_id)
                or (file_mappings or {}).get(business_name)
                or []
            )
        )
        mapped_folder = folders_by_id.get(mapped_folder_id) if mapped_folder_id else None
        mapped_folder_is_valid = bool(
            mapped_folder
            and mapped_folder_id in portfolio_descendants
            and mapped_folder_id != portfolio_folder_id
        )
        if mapped_folder_id and not mapped_folder_is_valid:
            matches[business_id] = {
                "status": "mapping_invalid",
                "business_name": business_name,
                "files": [],
                "folder_candidates": [
                    folder_details(folder) for folder in candidates
                ],
                "error": (
                    "The configured brand folder ID is not a folder inside the "
                    "verified portfolio folder."
                ),
            }
            continue
        if (
            not mapped_folder
            and len(candidates) != 1
            and (not mapped_file_ids or candidates)
        ):
            candidate_folders = list(candidates)
            if not candidate_folders and business_key:
                close_keys = difflib.get_close_matches(
                    business_key,
                    folder_index,
                    n=5,
                    cutoff=0.55,
                )
                candidate_folders = [
                    folder
                    for close_key in close_keys
                    for folder in folder_index[close_key]
                ]
            matches[business_id] = {
                "status": "ambiguous" if candidates else "unmatched",
                "business_name": business.get("name", ""),
                "files": [],
                "folder_candidates": [
                    folder_details(folder) for folder in candidate_folders
                ],
            }
            continue

        mapped_files = []
        for file_id in mapped_file_ids:
            try:
                file_item = (
                    service.files()
                    .get(
                        fileId=file_id,
                        fields=(
                            "id,name,mimeType,createdTime,modifiedTime,description,"
                            "webViewLink,parents"
                        ),
                        supportsAllDrives=True,
                    )
                    .execute(num_retries=3)
                )
            except (HttpError, httplib2.HttpLib2Error, OSError) as exc:
                logger.warning(
                    "Drive mapped file lookup failed business_id=%s file_id=%s "
                    "http_status=%s",
                    business_id,
                    file_id,
                    getattr(getattr(exc, "resp", None), "status", None),
                )
                matches[business_id] = {
                    "status": "file_mapping_invalid",
                    "business_name": business_name,
                    "files": [],
                    "error": (
                        "A configured Drive file could not be accessed by the "
                        "connected account."
                    ),
                }
                break
            parents = file_item.get("parents", [])
            if (
                not file_item.get("id")
                or file_item.get("mimeType") == _FOLDER_MIME_TYPE
                or not any(parent_id in portfolio_descendants for parent_id in parents)
            ):
                matches[business_id] = {
                    "status": "file_mapping_invalid",
                    "business_name": business_name,
                    "files": [],
                    "error": (
                        "A configured Drive file is not a document inside the "
                        "verified portfolio folder."
                    ),
                }
                break
            mapped_files.append(
                {
                    "id": file_item["id"],
                    "name": file_item.get("name", "(untitled)"),
                    "mime_type": file_item.get("mimeType", ""),
                    "created_time": file_item.get("createdTime"),
                    "modified_time": file_item.get("modifiedTime"),
                    "description": file_item.get("description", ""),
                    "web_view_link": file_item.get("webViewLink")
                    or f"https://drive.google.com/open?id={quote(file_item['id'], safe='')}",
                    "parent_ids": parents,
                    "content_available": False,
                    "content_status": "linked",
                }
            )
        if matches.get(business_id, {}).get("status") == "file_mapping_invalid":
            continue

        root_folder = mapped_folder or (candidates[0] if candidates else None)
        files: list[dict[str, Any]] = list(mapped_files)
        visited_file_ids: set[str] = {item["id"] for item in mapped_files}
        pending_folders = [root_folder] if root_folder else []
        visited_folder_ids: set[str] = set()
        folder_errors = []
        while pending_folders:
            current_folder = pending_folders.pop()
            folder_id = current_folder["id"]
            if folder_id in visited_folder_ids:
                continue
            visited_folder_ids.add(folder_id)
            page_token = None
            while True:
                params = {
                    "q": f"trashed = false and '{folder_id}' in parents",
                    "pageSize": 1000,
                    "fields": (
                        "nextPageToken,files(id,name,mimeType,createdTime,"
                        "modifiedTime,description,webViewLink,parents)"
                    ),
                    "supportsAllDrives": True,
                    "includeItemsFromAllDrives": True,
                    "corpora": "allDrives",
                }
                if page_token:
                    params["pageToken"] = page_token
                try:
                    result = service.files().list(**params).execute(num_retries=3)
                except (HttpError, httplib2.HttpLib2Error, OSError) as exc:
                    status_code = getattr(
                        getattr(exc, "resp", None),
                        "status",
                        None,
                    )
                    logger.warning(
                        "Drive folder listing failed business_id=%s "
                        "folder_id=%s http_status=%s exception_type=%s",
                        business_id,
                        folder_id,
                        status_code,
                        type(exc).__name__,
                    )
                    folder_errors.append(
                        f"{current_folder.get('name') or 'Drive folder'} "
                        "could not be listed."
                    )
                    break
                for item in result.get("files", []):
                    if item.get("mimeType") == _FOLDER_MIME_TYPE:
                        pending_folders.append(item)
                        continue
                    file_id = item.get("id")
                    if not file_id or file_id in visited_file_ids:
                        continue
                    visited_file_ids.add(file_id)
                    file_data = {
                        "id": file_id,
                        "name": item.get("name", "(untitled)"),
                        "mime_type": item.get("mimeType", ""),
                        "created_time": item.get("createdTime"),
                        "modified_time": item.get("modifiedTime"),
                        "description": item.get("description", ""),
                        "web_view_link": item.get("webViewLink")
                        or f"https://drive.google.com/open?id={quote(file_id, safe='')}",
                        "parent_ids": item.get("parents", []),
                        "content_available": False,
                        "content_status": "linked",
                    }
                    files.append(file_data)
                page_token = result.get("nextPageToken")
                if not page_token:
                    break

        matches[business_id] = {
            "status": "incomplete" if folder_errors else "matched",
            "business_name": business.get("name", ""),
            "folder_id": root_folder["id"] if root_folder else None,
            "folder_name": root_folder.get("name", "") if root_folder else "",
            "folder_link": root_folder.get("webViewLink", "") if root_folder else "",
            "files": files,
            "error": "; ".join(folder_errors),
            "truncated": False,
        }
    return matches