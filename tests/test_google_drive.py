import io
import os
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

from googleapiclient.errors import HttpError

from app import agent, db, main
from app.brand_import import merge_drive_context
from app.config import Settings, settings
from app.integrations.google_drive import (
    _extract_downloaded_file,
    _spreadsheet_text,
    _read_file_content,
    fetch_brand_drive_files,
    fetch_recent_drive_files,
)


class GoogleDriveTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_data_dir = settings.data_dir
        settings.data_dir = self.temp_dir.name
        db.init_db()
        self.user = db.create_user("drive-test-user", "not-used")
        self.service = MagicMock()
        self.sheets_service = MagicMock()
        self.drive_files = self.service.files.return_value
        self.spreadsheets = self.sheets_service.spreadsheets.return_value
        self.drive_files.list.return_value.execute.return_value = {
            "files": [
                {
                    "id": "doc-1",
                    "name": "Project brief",
                    "mimeType": "application/vnd.google-apps.document",
                    "createdTime": "2026-01-02T00:00:00Z",
                    "modifiedTime": "2026-02-03T00:00:00Z",
                    "webViewLink": "https://drive.google.com/file/d/doc-1",
                    "parents": ["folder-1"],
                },
                {
                    "id": "sheet-1",
                    "name": "Forecast",
                    "mimeType": "application/vnd.google-apps.spreadsheet",
                    "parents": ["folder-1"],
                },
                {
                    "id": "text-1",
                    "name": "Notes",
                    "mimeType": "text/plain",
                },
                {
                    "id": "pdf-1",
                    "name": "Archive",
                    "mimeType": "application/pdf",
                },
                {
                    "id": "folder-1",
                    "name": "Projects",
                    "mimeType": "application/vnd.google-apps.folder",
                    "webViewLink": "https://drive.google.com/drive/folders/folder-1",
                },
            ]
        }
        self.drive_files.get.return_value.execute.return_value = {
            "id": "portfolio-root",
            "name": "Portfolio",
            "mimeType": "application/vnd.google-apps.folder",
            "webViewLink": "https://drive.google.com/drive/folders/portfolio-root",
        }
        self.drive_files.export.return_value.execute.return_value = (
            b"Readable project brief"
        )
        self.drive_files.get_media.return_value.execute.return_value = b"Plain text contents"
        self.spreadsheets.get.return_value.execute.return_value = {
            "sheets": [
                {
                    "properties": {
                        "title": "Forecast",
                        "gridProperties": {"rowCount": 20, "columnCount": 5},
                    }
                }
            ]
        }
        self.spreadsheets.values.return_value.batchGet.return_value.execute.return_value = {
            "valueRanges": [
                {
                    "values": [
                        ["Month", "Revenue"],
                        ["January", "100"],
                    ]
                }
            ]
        }

    def tearDown(self):
        settings.data_dir = self.old_data_dir
        self.temp_dir.cleanup()

    def test_brand_folder_id_overrides_load_from_json_environment_setting(self):
        with patch.dict(
            os.environ,
            {"GOOGLE_DRIVE_BRAND_FOLDER_MAP": '{"Example Brand":"folder-id"}'},
        ):
            configured = Settings(_env_file=None)

        self.assertEqual(
            configured.google_drive_brand_folder_map,
            {"Example Brand": "folder-id"},
        )

    def test_brand_file_ids_load_from_json_environment_setting(self):
        with patch.dict(
            os.environ,
            {
                "GOOGLE_DRIVE_BRAND_FILE_MAP": (
                    '{"Cubsi":["cubsi-agreement"],'
                    '"Faction Labs":["faction-agreement"]}'
                )
            },
        ):
            configured = Settings(_env_file=None)

        self.assertEqual(
            configured.google_drive_brand_file_map,
            {
                "Cubsi": ["cubsi-agreement"],
                "Faction Labs": ["faction-agreement"],
            },
        )

    def _fetch(self, search_text="project revenue", max_results=5):
        credentials = object()
        with patch(
            "pypdf.PdfReader",
            return_value=SimpleNamespace(
                pages=[SimpleNamespace(extract_text=lambda: "Readable PDF text")]
            ),
        ), patch(
            "googleapiclient.discovery.build",
            side_effect=[self.service, self.sheets_service],
        ) as build, patch(
            "app.integrations.google_services.load_credentials",
            return_value=credentials,
        ) as load_credentials:
            files = fetch_recent_drive_files("user-1", search_text, max_results)
        self.assertEqual(build.call_count, 2)
        load_credentials.assert_called_once_with("user-1")
        self.assertTrue(
            all(call.kwargs["credentials"] is credentials for call in build.call_args_list)
        )
        return files

    def test_searches_and_reads_supported_types_and_returns_metadata(self):
        files = self._fetch()

        self.assertEqual(
            [file["content_status"] for file in files],
            ["read", "read", "read", "read", "folder"],
        )
        self.assertEqual(files[0]["content"], "Readable project brief")
        self.assertEqual(files[0]["created_time"], "2026-01-02T00:00:00Z")
        self.assertEqual(files[0]["parent_ids"], ["folder-1"])
        self.assertEqual(files[0]["parents"][0]["name"], "Portfolio")
        self.assertEqual(files[0]["web_view_link"], "https://drive.google.com/file/d/doc-1")
        self.assertEqual(
            files[1]["content"],
            "Sheet: Forecast\nMonth\tRevenue\nJanuary\t100",
        )
        self.assertEqual(files[1]["content_format"], "text/plain")
        requested_ranges = self.spreadsheets.values.return_value.batchGet.call_args.kwargs[
            "ranges"
        ]
        self.assertIn("'Forecast'!A1:E20", requested_ranges)
        self.drive_files.export.assert_called_once_with(
            fileId="doc-1",
            mimeType="text/plain",
        )
        self.assertEqual(files[2]["content"], "Plain text contents")
        self.assertTrue(files[3]["content_available"])
        self.assertEqual(files[3]["content"], "Readable PDF text")
        self.assertTrue(files[4]["is_folder"])
        self.assertFalse(files[4]["content_available"])
        self.assertNotIn("content_error", files[4])
        self.assertEqual(files[4]["name"], "Projects")

        query = self.drive_files.list.call_args.kwargs["q"]
        self.assertIn("fullText contains 'project' or fullText contains 'revenue'", query)
        self.assertNotIn("mimeType != 'application/vnd.google-apps.folder'", query)
        self.assertTrue(self.drive_files.list.call_args.kwargs["supportsAllDrives"])

    def test_cubsi_search_variants_use_the_same_drive_query(self):
        queries = []
        for search_text in ("Cubsi", "CUBSY", "cubsi"):
            self._fetch(search_text=search_text)
            queries.append(self.drive_files.list.call_args.kwargs["q"])

        self.assertEqual(queries[0], queries[1])
        self.assertEqual(queries[1], queries[2])
        self.assertIn("fullText contains 'Cubsi'", queries[0])
        self.assertIn("fullText contains 'CUBSY'", queries[0])

    def test_individual_content_error_does_not_hide_other_files(self):
        denied = HttpError(
            SimpleNamespace(
                status=403,
                reason=(
                    "access denied refresh_token=do-not-log "
                    "Bearer do-not-log-this-either"
                ),
            ),
            b'{"error":{"message":"access denied"}}',
        )
        self.spreadsheets.values.return_value.batchGet.return_value.execute.side_effect = denied

        with self.assertLogs("app.integrations.google_drive", level="WARNING") as logs:
            files = self._fetch()

        self.assertEqual(files[1]["content_status"], "error")
        self.assertFalse(files[1]["content_available"])
        self.assertEqual(
            files[1]["content_error"],
            "The connected Google account does not have permission to read this file.",
        )
        self.assertEqual(files[0]["content"], "Readable project brief")
        self.assertEqual(files[2]["content"], "Plain text contents")
        self.assertIn("file_name='Forecast'", logs.output[0])
        self.assertIn("file_id=sheet-1", logs.output[0])
        self.assertIn("mime_type=application/vnd.google-apps.spreadsheet", logs.output[0])
        self.assertIn("http_status=403", logs.output[0])
        self.assertIn("google_api_error=", logs.output[0])
        self.assertNotIn("do-not-log", logs.output[0])
        self.assertNotIn("do-not-log-this-either", logs.output[0])
        self.assertIn(
            "sheets.spreadsheets.get -> sheets.spreadsheets.values.batchGet",
            logs.output[0],
        )

    def test_missing_parser_is_actionable_and_does_not_abort_search(self):
        self.drive_files.list.return_value.execute.return_value = {
            "files": [
                {"id": "pdf-missing-parser", "name": "Agreement", "mimeType": "application/pdf"},
                {"id": "text-readable", "name": "Notes", "mimeType": "text/plain"},
            ]
        }
        self.drive_files.get_media.return_value.execute.side_effect = [
            b"PDF bytes",
            b"Readable notes",
        ]
        credentials = object()

        with patch.dict("sys.modules", {"pypdf": None}), patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=credentials,
        ), self.assertLogs("app.integrations.google_drive", level="WARNING") as logs:
            files = fetch_recent_drive_files("user-1", "agreement", 5)

        self.assertEqual(len(files), 2)
        self.assertEqual(files[0]["content_status"], "error")
        self.assertIn("pypdf", files[0]["content_error"])
        self.assertIn("pip install -r requirements.txt", files[0]["content_error"])
        self.assertEqual(files[1]["content_status"], "read")
        self.assertEqual(files[1]["content"], "Readable notes")
        self.assertIn("exception_type=ModuleNotFoundError", logs.output[0])

    def test_brand_files_are_matched_by_folder_name_and_read_recursively(self):
        nested_folder = {
            "id": "folder-nested",
            "name": "Archive",
            "mimeType": "application/vnd.google-apps.folder",
        }
        doc = {
            "id": "doc-1",
            "name": "Brand brief",
            "mimeType": "application/vnd.google-apps.document",
            "webViewLink": "https://drive.google.com/file/d/doc-1",
        }
        nested_doc = {
            "id": "doc-2",
            "name": "Old notes",
            "mimeType": "text/plain",
        }
        responses = [
            {
                "files": [
                    {
                        "id": "brand-1",
                        "name": "Example Brand",
                        "parents": ["portfolio-root"],
                    },
                    {
                        "id": "outside-brand",
                        "name": "Example Brand",
                        "parents": ["another-root"],
                    },
                ]
            },
            {"files": [nested_folder, doc]},
            {"files": [nested_doc]},
        ]
        self.drive_files.list.side_effect = [
            SimpleNamespace(execute=MagicMock(return_value=result))
            for result in responses
        ]
        businesses = [
            {"id": "business-1", "name": "Example-Brand"},
            {"id": "business-2", "name": "No Drive Folder"},
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                businesses,
                portfolio_folder_id="portfolio-root",
            )

        self.assertEqual(result["business-1"]["status"], "matched")
        self.assertEqual(
            [file["name"] for file in result["business-1"]["files"]],
            ["Brand brief", "Old notes"],
        )
        self.assertEqual(
            [file["content_status"] for file in result["business-1"]["files"]],
            ["linked", "linked"],
        )
        self.assertEqual(
            result["business-1"]["files"][0]["web_view_link"],
            "https://drive.google.com/file/d/doc-1",
        )
        self.assertEqual(
            result["business-1"]["files"][1]["web_view_link"],
            "https://drive.google.com/open?id=doc-2",
        )
        self.drive_files.get_media.assert_not_called()
        self.drive_files.export.assert_not_called()
        self.assertEqual(result["business-2"]["status"], "unmatched")
        queries = [call.kwargs["q"] for call in self.drive_files.list.call_args_list]
        self.assertIn("'brand-1' in parents", queries[1])
        self.assertIn("'folder-nested' in parents", queries[2])

    def test_portfolio_root_is_required_before_brand_matching(self):
        with patch("googleapiclient.discovery.build") as build, patch(
            "app.integrations.google_services.load_credentials"
        ) as load_credentials:
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-1", "name": "Example Brand"}],
            )

        self.assertEqual(
            result["business-1"]["status"],
            "portfolio_not_configured",
        )
        self.assertIn(
            "GOOGLE_DRIVE_PORTFOLIO_FOLDER_ID",
            result["business-1"]["error"],
        )
        build.assert_not_called()
        load_credentials.assert_not_called()

    def test_duplicate_brand_folders_are_marked_ambiguous(self):
        self.drive_files.list.return_value.execute.return_value = {
            "files": [
                {
                    "id": "folder-1",
                    "name": "Example Brand",
                    "parents": ["portfolio-root"],
                },
                {
                    "id": "folder-2",
                    "name": "Example Brand",
                    "parents": ["portfolio-root"],
                },
            ]
        }

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-1", "name": "Example Brand"}],
                portfolio_folder_id="portfolio-root",
            )

        self.assertEqual(result["business-1"]["status"], "ambiguous")
        self.assertEqual(
            {folder["id"] for folder in result["business-1"]["folder_candidates"]},
            {"folder-1", "folder-2"},
        )
        self.assertEqual(self.drive_files.list.call_count, 1)

    def test_explicit_folder_mapping_resolves_duplicate_brand_folders(self):
        responses = [
            {
                "files": [
                    {
                        "id": "folder-1",
                        "name": "Example Brand",
                        "mimeType": "application/vnd.google-apps.folder",
                        "parents": ["portfolio-root"],
                    },
                    {
                        "id": "folder-2",
                        "name": "Example Brand",
                        "mimeType": "application/vnd.google-apps.folder",
                        "parents": ["portfolio-root"],
                    },
                ]
            },
            {"files": []},
        ]
        self.drive_files.list.side_effect = [
            SimpleNamespace(execute=MagicMock(return_value=result))
            for result in responses
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-1", "name": "Example Brand"}],
                portfolio_folder_id="portfolio-root",
                folder_mappings={"Example Brand": "folder-2"},
            )

        self.assertEqual(result["business-1"]["status"], "matched")
        self.assertEqual(result["business-1"]["folder_id"], "folder-2")
        self.assertEqual(self.drive_files.list.call_count, 2)

    def test_explicit_file_mapping_links_documents_without_a_brand_folder(self):
        self.drive_files.list.return_value.execute.return_value = {"files": []}
        self.drive_files.get.side_effect = [
            SimpleNamespace(
                execute=MagicMock(
                    return_value={
                        "id": "portfolio-root",
                        "name": "Portfolio",
                        "mimeType": "application/vnd.google-apps.folder",
                    }
                )
            ),
            SimpleNamespace(
                execute=MagicMock(
                    return_value={
                        "id": "cubsi-agreement",
                        "name": "CUBSY X NCC AGREEMENT.pdf",
                        "mimeType": "application/pdf",
                        "parents": ["portfolio-root"],
                    }
                )
            ),
            SimpleNamespace(
                execute=MagicMock(
                    return_value={
                        "id": "faction-agreement",
                        "name": "NCC_X_Faction_agreement.pdf",
                        "mimeType": "application/pdf",
                        "parents": ["portfolio-root"],
                    }
                )
            ),
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [
                    {"id": "business-cubsi", "name": "Cubsi"},
                    {"id": "business-faction", "name": "Faction Labs"},
                ],
                portfolio_folder_id="portfolio-root",
                file_mappings={
                    "Cubsi": ["cubsi-agreement"],
                    "Faction Labs": ["faction-agreement"],
                },
            )

        self.assertEqual(result["business-cubsi"]["status"], "matched")
        self.assertEqual(
            [file["name"] for file in result["business-cubsi"]["files"]],
            ["CUBSY X NCC AGREEMENT.pdf"],
        )
        self.assertEqual(
            result["business-cubsi"]["files"][0]["content_status"],
            "linked",
        )
        self.assertIsNone(result["business-cubsi"]["folder_id"])
        self.assertEqual(result["business-faction"]["status"], "matched")
        self.assertEqual(
            [file["name"] for file in result["business-faction"]["files"]],
            ["NCC_X_Faction_agreement.pdf"],
        )
        self.drive_files.get_media.assert_not_called()
        self.drive_files.export.assert_not_called()

    def test_explicit_file_mapping_rejects_a_file_outside_portfolio(self):
        self.drive_files.list.return_value.execute.return_value = {"files": []}
        self.drive_files.get.side_effect = [
            SimpleNamespace(
                execute=MagicMock(
                    return_value={
                        "id": "portfolio-root",
                        "name": "Portfolio",
                        "mimeType": "application/vnd.google-apps.folder",
                    }
                )
            ),
            SimpleNamespace(
                execute=MagicMock(
                    return_value={
                        "id": "agreement-outside-portfolio",
                        "name": "Unrelated agreement.pdf",
                        "mimeType": "application/pdf",
                        "parents": ["other-folder"],
                    }
                )
            ),
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-cubsi", "name": "Cubsi"}],
                portfolio_folder_id="portfolio-root",
                file_mappings={
                    "Cubsi": ["agreement-outside-portfolio"],
                },
            )

        self.assertEqual(result["business-cubsi"]["status"], "file_mapping_invalid")
        self.assertEqual(result["business-cubsi"]["files"], [])
        self.assertIn("inside the verified portfolio", result["business-cubsi"]["error"])

    def test_brand_files_over_one_hundred_are_not_skipped(self):
        root = {
            "id": "brand-folder",
            "name": "Example Brand",
            "mimeType": "application/vnd.google-apps.folder",
            "parents": ["portfolio-root"],
        }
        files = [
            {
                "id": f"file-{index}",
                "name": f"Notes {index}",
                "mimeType": "text/plain",
            }
            for index in range(101)
        ]
        self.drive_files.list.side_effect = [
            SimpleNamespace(execute=MagicMock(return_value={"files": [root]})),
            SimpleNamespace(execute=MagicMock(return_value={"files": files})),
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-1", "name": "Example Brand"}],
                portfolio_folder_id="portfolio-root",
            )

        self.assertEqual(len(result["business-1"]["files"]), 101)
        self.assertFalse(result["business-1"]["truncated"])

    def test_uploaded_pdf_office_csv_and_html_content_is_extracted(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches

        doc = Document()
        doc.add_paragraph("Agreement payment terms")
        docx_content = io.BytesIO()
        doc.save(docx_content)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Pricing"
        worksheet.append(["Product", "Price"])
        worksheet.append(["Sample", 12.5])
        xlsx_content = io.BytesIO()
        workbook.save(xlsx_content)

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(1),
        )
        textbox.text_frame.text = "Launch plan"
        pptx_content = io.BytesIO()
        presentation.save(pptx_content)

        test_files = [
            (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                docx_content.getvalue(),
                "Agreement payment terms",
            ),
            (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                xlsx_content.getvalue(),
                "Sample\t12.5",
            ),
            (
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                pptx_content.getvalue(),
                "Launch plan",
            ),
            ("text/csv", b"Product,Price\nSample,12.5", "Sample\t12.5"),
            (
                "text/html",
                b"<html><script>ignore()</script><p>Product description</p></html>",
                "Product description",
            ),
        ]
        for mime_type, content, expected in test_files:
            with self.subTest(mime_type=mime_type):
                self.drive_files.get_media.return_value.execute.return_value = content
                result = _read_file_content(
                    self.service,
                    object(),
                    {"id": "uploaded-file", "mimeType": mime_type},
                )
                self.assertEqual(result["content_status"], "read")
                self.assertIn(expected, result["content"])

        with patch(
            "pypdf.PdfReader",
            return_value=SimpleNamespace(
                pages=[SimpleNamespace(extract_text=lambda: "Signed agreement terms")]
            ),
        ):
            self.drive_files.get_media.return_value.execute.return_value = b"PDF bytes"
            pdf_result = _read_file_content(
                self.service,
                object(),
                {"id": "pdf-file", "mimeType": "application/pdf"},
            )
        self.assertEqual(pdf_result["content_status"], "read")
        self.assertIn("Signed agreement terms", pdf_result["content"])

    def test_readable_pdf_is_parsed_with_pdf_reader(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=100)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 50 Td (Readable agreement terms) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        pdf_content = io.BytesIO()
        writer.write(pdf_content)

        text, truncated = _extract_downloaded_file(
            pdf_content.getvalue(),
            "application/pdf",
        )

        self.assertIn("Readable agreement terms", text)
        self.assertFalse(truncated)

    def test_missing_spreadsheet_parser_returns_actionable_error(self):
        with patch.dict("sys.modules", {"openpyxl": None}):
            result = _read_file_content(
                self.service,
                object(),
                {
                    "id": "xlsx-file",
                    "name": "Agreement terms",
                    "mimeType": (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                },
            )

        self.assertEqual(result["content_status"], "error")
        self.assertIn("openpyxl", result["content_error"])
        self.assertIn("pip install -r requirements.txt", result["content_error"])

    def test_malformed_and_empty_files_return_read_errors(self):
        for content, mime_type in (
            (b"not a PDF", "application/pdf"),
            (b"", "text/plain"),
        ):
            with self.subTest(mime_type=mime_type), self.assertLogs(
                "app.integrations.google_drive",
                level="WARNING",
            ):
                self.drive_files.get_media.return_value.execute.return_value = content
                result = _read_file_content(
                    self.service,
                    object(),
                    {
                        "id": "unreadable-file",
                        "name": "Unreadable",
                        "mimeType": mime_type,
                    },
                )

            self.assertEqual(result["content_status"], "error")
            self.assertFalse(result["content_available"])
            self.assertTrue(result["content_error"])

    def test_brand_folder_and_file_listing_follow_pagination_tokens(self):
        root = {
            "id": "brand-folder",
            "name": "Example Brand",
            "mimeType": "application/vnd.google-apps.folder",
            "parents": ["portfolio-root"],
        }
        responses = [
            {"files": [root], "nextPageToken": "folders-next"},
            {"files": []},
            {
                "files": [
                    {
                        "id": "doc-1",
                        "name": "First",
                        "mimeType": "text/plain",
                    }
                ],
                "nextPageToken": "files-next",
            },
            {
                "files": [
                    {
                        "id": "doc-2",
                        "name": "Second",
                        "mimeType": "text/plain",
                    }
                ]
            },
        ]
        self.drive_files.list.side_effect = [
            SimpleNamespace(execute=MagicMock(return_value=result))
            for result in responses
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-1", "name": "Example Brand"}],
                portfolio_folder_id="portfolio-root",
            )

        self.assertEqual(
            [file["id"] for file in result["business-1"]["files"]],
            ["doc-1", "doc-2"],
        )
        page_tokens = [
            call.kwargs.get("pageToken")
            for call in self.drive_files.list.call_args_list
        ]
        self.assertIn("folders-next", page_tokens)
        self.assertIn("files-next", page_tokens)

    def test_brand_folder_sync_links_files_without_fetching_contents(self):
        responses = [
            {
                "files": [
                    {
                        "id": "brand-folder",
                        "name": "Example Brand",
                        "mimeType": "application/vnd.google-apps.folder",
                        "parents": ["portfolio-root"],
                    }
                ]
            },
            {
                "files": [
                    {
                        "id": "doc-1",
                        "name": "Brief",
                        "mimeType": "text/plain",
                        "modifiedTime": "2026-02-03T00:00:00Z",
                    }
                ]
            },
        ]
        self.drive_files.list.side_effect = [
            SimpleNamespace(execute=MagicMock(return_value=result))
            for result in responses
        ]

        with patch(
            "googleapiclient.discovery.build",
            return_value=self.service,
        ), patch(
            "app.integrations.google_services.load_credentials",
            return_value=object(),
        ):
            result = fetch_brand_drive_files(
                "user-1",
                [{"id": "business-1", "name": "Example Brand"}],
                portfolio_folder_id="portfolio-root",
            )

        file = result["business-1"]["files"][0]
        self.assertEqual(file["content_status"], "linked")
        self.assertEqual(
            file["web_view_link"],
            "https://drive.google.com/open?id=doc-1",
        )
        self.assertNotIn("content", file)
        self.drive_files.get_media.assert_not_called()
        self.drive_files.export.assert_not_called()

    def test_google_slides_text_is_extracted(self):
        slides_service = MagicMock()
        slides_service.presentations.return_value.get.return_value.execute.return_value = {
            "slides": [
                {
                    "pageElements": [
                        {
                            "shape": {
                                "text": {
                                    "textElements": [
                                        {"textRun": {"content": "Launch plan"}},
                                        {"textRun": {"content": "\nTarget: Q4"}},
                                    ]
                                }
                            }
                        }
                    ]
                }
            ]
        }

        with patch(
            "googleapiclient.discovery.build",
            return_value=slides_service,
        ):
            result = _read_file_content(
                self.service,
                object(),
                {
                    "id": "slide-1",
                    "mimeType": "application/vnd.google-apps.presentation",
                },
            )

        self.assertEqual(result["content_status"], "read")
        self.assertEqual(result["content"], "Launch plan\n\nTarget: Q4")
        slides_service.presentations.return_value.get.assert_called_once()

    def test_spreadsheet_row_and_column_caps_are_reported_as_truncated(self):
        sheets = MagicMock()
        sheets.spreadsheets.return_value.get.return_value.execute.return_value = {
            "sheets": [
                {
                    "properties": {
                        "title": "Large",
                        "gridProperties": {"rowCount": 501, "columnCount": 51},
                    }
                }
            ]
        }
        sheets.spreadsheets.return_value.values.return_value.batchGet.return_value.execute.return_value = {
            "valueRanges": [{"values": [["Header"]]}]
        }

        content, truncated = _spreadsheet_text(sheets, "sheet-1", [])

        self.assertIn("Sheet: Large", content)
        self.assertTrue(truncated)
        requested_range = (
            sheets.spreadsheets.return_value.values.return_value.batchGet.call_args
            .kwargs["ranges"][0]
        )
        self.assertEqual(requested_range, "'Large'!A1:AX500")

    def test_drive_context_merge_preserves_manual_context_and_replaces_old_import(self):
        original = (
            "# Example Brand\n\nManual notes.\n\n"
            "<!-- google-drive-import:start -->\nOld imported text\n"
            "<!-- google-drive-import:end -->"
        )
        files = [
            {
                "name": "Brand brief",
                "content_status": "linked",
                "web_view_link": "https://drive.google.com/file/d/doc-1",
            }
        ]

        merged = merge_drive_context(original, files)
        merged_again = merge_drive_context(merged, files)

        self.assertIn("Manual notes.", merged)
        self.assertIn(
            "[Open in Google Drive](https://drive.google.com/file/d/doc-1)",
            merged,
        )
        self.assertNotIn("Content not extracted", merged)
        self.assertNotIn("Old imported text", merged)
        self.assertEqual(merged_again, merged)

    def test_agent_passes_drive_content_to_the_existing_model_turn(self):
        tool_call = SimpleNamespace(
            id="call-drive",
            function=SimpleNamespace(
                name="search_google_drive",
                arguments='{"search_text":"forecast"}',
            ),
        )
        client = MagicMock()
        client.chat.completions.create.side_effect = [
            SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(content=None, tool_calls=[tool_call])
                    )
                ]
            ),
            SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(
                            content="The forecast shows January revenue of 100.",
                            tool_calls=None,
                        )
                    )
                ]
            ),
        ]
        drive_data = [
            {
                "name": "Forecast",
                "content": "Month,Revenue\nJanuary,100",
                "content_status": "read",
            }
        ]

        with patch("app.agent._client", return_value=client), patch(
            "app.agent.db.attention_brief", return_value={"summary": ""}
        ), patch("app.agent.db.list_businesses", return_value=[]), patch(
            "app.integrations.google_drive.fetch_recent_drive_files",
            return_value=drive_data,
        ) as fetch_drive:
            result = agent.chat("What does the forecast say?", user_id="user-1", persist=False)

        self.assertIn("January revenue of 100", result["reply"])
        fetch_drive.assert_called_once_with(
            "user-1", search_text="forecast", max_results=5
        )
        tool_message = client.chat.completions.create.call_args_list[1].kwargs[
            "messages"
        ][-1]
        self.assertEqual(tool_message["role"], "tool")
        self.assertIn("Month,Revenue", tool_message["content"])

    def test_drive_pull_endpoint_records_status_and_returns_files(self):
        db.save_google_token(
            self.user["id"],
            {
                "token": "access-token",
                "refresh_token": "refresh-token",
                "scopes": [],
            },
        )
        pulled_files = [
            {"id": "doc-1", "name": "Project brief", "content_status": "read"}
        ]

        with patch(
            "app.main.auth.require_super_admin",
            return_value=self.user,
        ), patch(
            "app.integrations.google_drive.fetch_recent_drive_files",
            return_value=pulled_files,
        ) as fetch_files:
            response = main.google_drive(request=object(), max_results=100)

        self.assertEqual(response, {"files": pulled_files, "count": 1})
        fetch_files.assert_called_once_with(
            self.user["id"],
            max_results=10,
        )
        token = db.get_google_token(self.user["id"])
        self.assertIsNotNone(token["last_drive_pull_at"])
        self.assertEqual(token["last_drive_count"], 1)
        db.save_google_token(
            self.user["id"],
            {"token": "refreshed-access-token", "scopes": []},
        )
        token = db.get_google_token(self.user["id"])
        self.assertIsNotNone(token["last_drive_pull_at"])
        self.assertEqual(token["last_drive_count"], 1)

    def test_drive_brand_import_updates_only_matched_contexts(self):
        business = db.create_business(
            name="Example Brand",
            description="",
            context_md="# Example Brand\n\nManual notes.",
        )
        db.save_google_token(
            self.user["id"],
            {"token": "access-token", "refresh_token": "refresh-token", "scopes": []},
        )
        fetched = {
            business["id"]: {
                "status": "matched",
                "folder_name": "Example Brand",
                "files": [
                    {
                        "id": "doc-1",
                        "name": "Brand brief",
                        "mime_type": "application/vnd.google-apps.document",
                        "content_status": "linked",
                        "web_view_link": "https://drive.google.com/file/d/doc-1",
                    }
                ],
            }
        }

        with patch(
            "app.main.auth.require_super_admin",
            return_value=self.user,
        ), patch(
            "app.integrations.google_drive.verify_drive_access",
            return_value={"drive_status": "verified", "drive_error": ""},
        ), patch(
            "app.integrations.google_drive.fetch_brand_drive_files",
            return_value=fetched,
        ) as fetch_files:
            response = main.import_google_drive_brands(request=object())
            unchanged_response = main.import_google_drive_brands(request=object())

        self.assertEqual(fetch_files.call_count, 2)
        self.assertEqual(
            fetch_files.call_args_list[0].args,
            (self.user["id"], [business]),
        )
        self.assertEqual(
            fetch_files.call_args_list[0].kwargs["portfolio_folder_id"],
            settings.google_drive_portfolio_folder_id,
        )
        self.assertEqual(
            fetch_files.call_args_list[0].kwargs["folder_mappings"],
            settings.google_drive_brand_folder_map,
        )
        updated = db.get_business(business["id"])
        self.assertIn("Manual notes.", updated["context_md"])
        self.assertIn("Drive file ID: doc-1", updated["context_md"])
        self.assertIn(
            "[Open in Google Drive](https://drive.google.com/file/d/doc-1)",
            updated["context_md"],
        )
        self.assertEqual(response["updated_count"], 1)
        self.assertEqual(response["brands"][0]["linked_count"], 1)
        self.assertEqual(unchanged_response["updated_count"], 0)
        self.assertEqual(unchanged_response["brands"][0]["status"], "unchanged")
        with db.connect() as conn:
            sync = conn.execute(
                "SELECT status, files_seen, files_read, last_successful_at "
                "FROM google_drive_brand_sync "
                "WHERE user_id = ? AND business_id = ?",
                (self.user["id"], business["id"]),
            ).fetchone()
            sources = conn.execute(
                "SELECT file_id, name, content_status FROM google_drive_sources "
                "WHERE user_id = ? AND business_id = ?",
                (self.user["id"], business["id"]),
            ).fetchall()
        self.assertEqual(tuple(sync)[:3], ("unchanged", 1, 0))
        self.assertIsNotNone(sync["last_successful_at"])
        self.assertEqual(
            [tuple(source) for source in sources],
            [("doc-1", "Brand brief", "linked")],
        )

    def test_link_only_brand_sync_imports_supported_and_unsupported_file_links(self):
        business = db.create_business(
            name="Example Brand",
            description="",
            context_md="# Example Brand\n\nPreviously verified content.",
        )
        db.save_google_token(
            self.user["id"],
            {"token": "access-token", "refresh_token": "refresh-token", "scopes": []},
        )
        fetched = {
            business["id"]: {
                "status": "matched",
                "folder_id": "folder-1",
                "folder_name": "Example Brand",
                "files": [
                    {
                        "id": "doc-1",
                        "name": "New brief",
                        "mime_type": "application/vnd.google-apps.document",
                        "web_view_link": "https://drive.google.com/file/d/doc-1",
                        "content_status": "linked",
                    },
                    {
                        "id": "pdf-1",
                        "name": "Terms.pdf",
                        "mime_type": "application/x-iwork-pages-sffpages",
                        "web_view_link": "https://drive.google.com/file/d/pdf-1",
                        "content_status": "linked",
                    },
                ],
            }
        }

        with patch(
            "app.main.auth.require_super_admin",
            return_value=self.user,
        ), patch(
            "app.integrations.google_drive.verify_drive_access",
            return_value={"drive_status": "verified", "drive_error": ""},
        ), patch(
            "app.integrations.google_drive.fetch_brand_drive_files",
            return_value=fetched,
        ):
            response = main.import_google_drive_brands(request=object())

        updated = db.get_business(business["id"])
        self.assertEqual(response["sync_mode"], "links_only")
        self.assertIn(
            "[Open in Google Drive](https://drive.google.com/file/d/doc-1)",
            updated["context_md"],
        )
        self.assertIn("Terms.pdf", updated["context_md"])
        self.assertIn(
            "[Open in Google Drive](https://drive.google.com/file/d/pdf-1)",
            updated["context_md"],
        )
        self.assertEqual(response["brands"][0]["status"], "imported")
        self.assertEqual(response["brands"][0]["linked_count"], 2)
        self.assertEqual(response["brands"][0]["errors"], [])
        with db.connect() as conn:
            sync = conn.execute(
                "SELECT status, files_seen, files_read, error_count, last_successful_at "
                "FROM google_drive_brand_sync WHERE user_id = ? AND business_id = ?",
                (self.user["id"], business["id"]),
            ).fetchone()
            source_count = conn.execute(
                "SELECT COUNT(*) FROM google_drive_sources WHERE user_id = ?",
                (self.user["id"],),
            ).fetchone()[0]
        self.assertEqual(tuple(sync)[:4], ("imported", 2, 0, 0))
        self.assertIsNotNone(sync["last_successful_at"])
        self.assertEqual(source_count, 2)

    def test_links_only_brand_loads_inventory_without_reading_file_content(self):
        business = db.create_business(
            name="Example Brand",
            description="",
            context_md="# Example Brand\n\nExisting notes.",
        )
        db.save_google_token(
            self.user["id"],
            {"token": "access-token", "refresh_token": "refresh-token", "scopes": []},
        )
        fetched = {
            business["id"]: {
                "status": "matched",
                "folder_id": "folder-1",
                "folder_name": "Example Brand",
                "files": [
                    {
                        "id": "pdf-1",
                        "name": "Signed agreement.pdf",
                        "mime_type": "application/pdf",
                        "web_view_link": "https://drive.google.com/file/d/pdf-1",
                        "content_status": "linked",
                    }
                ],
            }
        }

        with patch(
            "app.main.auth.require_super_admin",
            return_value=self.user,
        ), patch(
            "app.integrations.google_drive.verify_drive_access",
            return_value={"drive_status": "verified", "drive_error": ""},
        ), patch(
            "app.integrations.google_drive.fetch_brand_drive_files",
            return_value=fetched,
        ):
            response = main.import_google_drive_brands(request=object())

        updated = db.get_business(business["id"])
        self.assertEqual(response["updated_count"], 1)
        self.assertEqual(response["brands"][0]["status"], "imported")
        self.assertIn("Existing notes.", updated["context_md"])
        self.assertIn("Signed agreement.pdf", updated["context_md"])
        self.assertIn(
            "[Open in Google Drive](https://drive.google.com/file/d/pdf-1)",
            updated["context_md"],
        )

    def test_brand_import_requires_verified_drive_access(self):
        business = db.create_business(name="Example Brand")

        with patch(
            "app.main.auth.require_super_admin",
            return_value=self.user,
        ), patch(
            "app.integrations.google_drive.verify_drive_access",
            return_value={
                "drive_status": "permission_missing",
                "drive_error": "Drive permission is missing.",
            },
        ), patch(
            "app.integrations.google_drive.fetch_brand_drive_files"
        ) as fetch_files:
            with self.assertRaises(main.HTTPException) as raised:
                main.import_google_drive_brands(request=object())

        self.assertEqual(raised.exception.status_code, 403)
        self.assertEqual(raised.exception.detail, "Drive permission is missing.")
        fetch_files.assert_not_called()
        sync = db.get_google_drive_business_sync(
            self.user["id"],
            business["id"],
        )
        self.assertEqual(sync["status"], "failed")
        self.assertEqual(sync["error"], "Drive permission is missing.")


if __name__ == "__main__":
    unittest.main()
